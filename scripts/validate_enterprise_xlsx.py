#!/usr/bin/env python3
"""Independent OOXML/openpyxl checks of the synthetic enterprise workbook.
FF_ENTERPRISE_FIXTURE_DIR=/tmp/enterprise swift test --filter EnterpriseGeotechnicalTests
python3 scripts/validate_enterprise_xlsx.py /tmp/enterprise
"""
import json
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
import openpyxl

root = Path(sys.argv[1])
expected = json.loads((root / 'expected.json').read_text())
path = root / 'FieldForge-Geotechnical.xlsx'
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(path) as archive:
    assert archive.testzip() is None
    assert len(archive.namelist()) == 10
    for name in archive.namelist():
        part = ET.fromstring(archive.read(name))
        assert not part.findall('.//s:f', ns), 'Unexpected formula'
book = openpyxl.load_workbook(path, data_only=False)
assert book.sheetnames == ['Геотех картирование', 'Справочники', 'Исходные данные']
assert all(s.sheet_state == 'visible' for s in book)
report, refs, raw = book.worksheets
assert report.max_row == expected['rows'] + 1
headers = {cell.value: cell.column for cell in report[1]}
assert abs(report.cell(2, headers['SD среднее, м']).value - expected['mean']) < 1e-15
assert report.cell(2, headers['SD среднее, м']).data_type == 'n'
assert report.cell(2, headers['Систем в среднем SD']).value == 3
assert report.cell(2, headers['SD содержит оценки по RQD']).value == 'Да'
assert report.cell(2, headers['ROCK · код']).value is None
assert report.cell(2, headers['J5 SD, м']).value is None
assert report.cell(2, headers['Примечания']).value == 'Кириллица =1+2'
assert report.cell(2, headers['Примечания']).data_type == 's'
assert report.cell(2, headers['Дата']).is_date
assert all(s.freeze_panes == 'A2' and s.auto_filter.ref for s in book)
values = {c.value for row in raw for c in row if c.value is not None}
assert {'manual', 'estimatedFromRQD', 'Исправление', 'Наблюдение'} <= values
assert '/ratings/gsi/selection/x' in values
assert '/snapshot/ratings/gsi/selection/x' in values
assert expected['estimated_sd'] in values  # exact round-trip SD input, distinct from Excel presentation
print(f'PASS: 3 visible sheets, {report.max_row-1} observation, {raw.max_row-1} scalar source rows; SD={expected["mean"]}')
