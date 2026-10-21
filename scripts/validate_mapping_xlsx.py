#!/usr/bin/env python3
"""Independent ZIP/OOXML + openpyxl validation of synthetic mapping workbooks.

P52_EXPORT_FIXTURE_DIR=/tmp/p52 swift test --filter MappingXLSXExportTests
python3 scripts/validate_mapping_xlsx.py /tmp/p52
FF_FINAL_EXPORT_FIXTURE_DIR=/tmp/final swift test --filter FinalFieldPolishTests
python3 scripts/validate_mapping_xlsx.py /tmp/final
Requires openpyxl for read-only validation; not an application dependency.
"""
import json
from datetime import datetime, timedelta
import posixpath
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
import openpyxl

root = Path(sys.argv[1])
expected = json.loads((root / 'expected.json').read_text())
path = root / 'FieldForge-Mapping.xlsx'
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(path) as archive:
    assert archive.testzip() is None, 'ZIP CRC mismatch'
    names = archive.namelist()
    assert len(names) == len(set(names))
    assert len(names) == len(parts_sheets := [n for n in names if n.startswith("xl/worksheets/")]) + 7
    assert not any('externalLink' in n or 'vba' in n for n in names)
    parts = {name: ET.fromstring(archive.read(name)) for name in names}
    for name, part in parts.items():
        assert not part.findall('.//s:f', ns), name
        if name.endswith('.rels'):
            base = '' if name == '_rels/.rels' else posixpath.dirname(posixpath.dirname(name))
            ids = set()
            for rel in part:
                assert rel.attrib['Id'] not in ids
                ids.add(rel.attrib['Id'])
                assert rel.attrib.get('TargetMode') != 'External'
                target = posixpath.normpath(posixpath.join(base, rel.attrib['Target']))
                assert target in names, (name, target)
    for child in parts['[Content_Types].xml']:
        if child.tag.endswith('Override'):
            assert child.attrib['PartName'].lstrip('/') in names
    for name in names:
        if name.startswith('xl/worksheets/'):
            sheet = parts[name]
            assert sheet.find('s:sheetViews/s:sheetView/s:pane', ns).attrib['state'] == 'frozen'
            assert sheet.find('s:mergeCells', ns) is None
            for cell in sheet.findall('.//s:c', ns):
                assert int(cell.attrib['s']) < 6
                assert cell.attrib.get('t') in (None, 'inlineStr', 'n')

book = openpyxl.load_workbook(path, data_only=False, keep_links=False)
field_names = ['Наблюдения', 'Системы трещин', 'Тектоника', 'Рейтинги массива', 'Фотографии']
raw_names = ['projects', 'horizons', 'excavations', 'excavation_horizons', 'plans', 'points', 'photo_point_contexts', 'measurements', 'joint_sets', 'interpretation_evidence', 'tectonic_structures', 'assessments', 'photos', 'photo_attachments']
assert book.sheetnames == field_names + ['Сводка'] + raw_names + ['field_links']
assert all(book[name].sheet_state == 'visible' for name in field_names + ['Сводка'])
assert all(book[name].sheet_state == 'hidden' for name in raw_names + ['field_links'])

checked = 0
for name, table in expected.items():
    sheet = book[name]
    assert sheet.freeze_panes == 'A2'
    assert sheet.max_row == len(table['rows']) + 1
    assert sheet.max_column == len(table['columns'])
    assert [c.value for c in sheet[1]] == table['columns']
    assert sheet.auto_filter.ref == sheet.calculate_dimension()
    assert not sheet.merged_cells
    assert sheet['A1'].font.bold and sheet['A1'].fill.fgColor.rgb == 'FF243746'
    for i, row in enumerate(table['rows'], 2):
        for j, expected_cell in enumerate(row, 1):
            cell = sheet.cell(i, j)
            if expected_cell['type'] == 'blank':
                assert cell.value is None, (name, cell.coordinate, cell.value)
            elif expected_cell['type'] == 'd':
                expected_date = datetime(1970, 1, 1) + timedelta(seconds=expected_cell['value'])
                assert cell.data_type == 'd' and abs((cell.value - expected_date).total_seconds()) < 0.0011, (name, cell.coordinate, cell.value, expected_date)
            else:
                assert cell.data_type == expected_cell['type'], (name, cell.coordinate, cell.data_type)
                assert cell.value == expected_cell['value'], (name, cell.coordinate, cell.value, expected_cell)
            checked += 1
        if name in raw_names and i % 2 == 1:
            assert sheet.cell(i, 1).fill.fgColor.rgb == 'FFF0F5F7'

def records(name):
    sheet = book[name]
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]

tables = {name: records(name) for name in expected}
def keys(name, column):
    values = [r[column] for r in tables[name]]
    assert len(values) == len(set(values)), (name, column, 'duplicate key')
    return set(values)

projects = keys('projects', 'project_id')
horizons = keys('horizons', 'horizon_id')
excavations = keys('excavations', 'excavation_id')
plans = keys('plans', 'plan_id')
points = keys('points', 'point_id')
measurements = keys('measurements', 'measurement_id')
sets = keys('joint_sets', 'joint_set_id')
attachments = keys('photo_attachments', 'attachment_id')
contexts = keys('photo_point_contexts', 'revision_id')
for table in ['horizons', 'excavations', 'points', 'photo_point_contexts', 'photos']:
    assert all(r['project_id'] in projects for r in tables[table])
for table in ['plans', 'points', 'photos', 'photo_point_contexts', 'excavation_horizons']:
    assert all(r['horizon_id'] in horizons for r in tables[table])
for table in ['points', 'photo_point_contexts', 'photos', 'excavation_horizons']:
    assert all(r['excavation_id'] in excavations for r in tables[table])
for table in ['measurements', 'joint_sets', 'tectonic_structures', 'assessments', 'photos', 'interpretation_evidence', 'photo_point_contexts']:
    assert all(r['point_id'] in points for r in tables[table])
for table in ['points', 'photo_point_contexts']:
    assert all(r['plan_id'] is None or r['plan_id'] in plans for r in tables[table])
membership = {(r['excavation_id'], r['horizon_id']) for r in tables['excavation_horizons']}
for table in ['points', 'photo_point_contexts', 'photos']:
    assert all((r['excavation_id'], r['horizon_id']) in membership for r in tables[table])
set_rows = {r['joint_set_id']: r for r in tables['joint_sets']}
measurement_rows = {r['measurement_id']: r for r in tables['measurements']}
context_rows = {r['revision_id']: r for r in tables['photo_point_contexts']}
plan_rows = {r['plan_id']: r for r in tables['plans']}
for table in ['points', 'photo_point_contexts']:
    assert all(r['plan_id'] is None or plan_rows[r['plan_id']]['horizon_id'] == r['horizon_id'] for r in tables[table])
assessment_keys = [(r['point_id'], r['kind']) for r in tables['assessments']]
assert len(assessment_keys) == len(set(assessment_keys))
for row in tables['interpretation_evidence']:
    assert row['measurement_id'] in measurements and row['joint_set_id'] in sets
    assert row['point_id'] == set_rows[row['joint_set_id']]['point_id'] == measurement_rows[row['measurement_id']]['point_id']
    assert row['joint_set_revision'] == set_rows[row['joint_set_id']]['revision']
for row in tables['photos']:
    assert row['attachment_id'] in attachments and row['context_revision_id'] in contexts
    context = context_rows[row['context_revision_id']]
    assert all(context[key] == row[key] for key in ['point_id', 'project_id', 'horizon_id', 'excavation_id'])
    assert context['version'] == row['context_point_version']
    assert row['captured_at_utc'] is None and row['capture_location_basis'] == 'unknownBeforeFirstSave'
    assert row['original_integrity'] == 'notChecked; originalsNotIncluded'
assert len(measurements) == 2
assert len(points) == (3 if 'Наблюдения' in expected else 1)
assert len(sets) == (5 if 'Наблюдения' in expected else 2)
assert any(r['x'] is not None for r in tables['points'])
assert all(book['points'].cell(i, 20).data_type == 's' for i, r in enumerate(tables['points'], 2) if r['x'] is not None)  # exact long decimal X
assert tables['photos'][0]['name'] in ('=1+1', '=1+1 Фото')
assert all(r['stored_q'] is None for r in tables['assessments'] if r['kind'] == 'RMR89-input-only')
if 'Наблюдения' in expected:
    observations = records('Наблюдения')
    assert len(observations) == 3
    assert {r['Проект'] for r in observations} == {'Рудник Северный'}
    assert {r['Выработка'] for r in observations} == {'Штрек Восточный № 3'}
    assert 'ПК 12+1,8' in {r['Пикет'] for r in observations}
    first = next(r for r in observations if r['RMR89 базовый'] is not None)
    assert first['RMR89 базовый'] == 64 and first['Класс RMR89 базовый'] == 'II — хорошее качество'
    assert first['X'] == 123456.12345678912 and first['Y'] == -0.123456789123456
    assert all(first[f'J{i} падение, °'] is not None for i in range(1, 6))
    assert all(symbol in first['Примечания'] for symbol in ['°', 'μ', 'µ', '≤'])
    assert any(r['X'] is None for r in observations)
    assert isinstance(first['Дата'], datetime)
    assert len({r['Дата'] for r in observations}) == 3
    assert book['Наблюдения']['G2'].number_format == '0.0'
    for r in records('Рейтинги массива'):
        if r['Метод'].startswith('Bieniawski-RMR89/basic'):
            assert r['RMR89 базовый'] == 64 and r['RMR балл состояния'] == 22
            assert [r[k] for k in ['RMR балл прочности', 'RMR балл RQD', 'RMR балл расстояния', 'RMR балл состояния', 'RMR балл воды']] == [12, 13, 10, 22, 7]
    links = records('field_links')
    refs = {r['Наблюдение'] for r in observations}
    point_links = {r['field_reference']: r for r in links if r['kind'] == 'observation'}
    assert set(point_links) == refs
    assert {r['point_id'] for r in point_links.values()} == points
    assert {r['record_id'] for r in links if r['kind'] == 'joint'} == sets
    assert all(r['Наблюдение'] in refs for r in records('Системы трещин') + records('Тектоника') + records('Фотографии'))
    assert all('Не включено' in r['Изображение в книге'] for r in records('Фотографии'))
    advanced = openpyxl.load_workbook(root / 'FieldForge-Technical.xlsx', data_only=False)
    assert advanced.sheetnames == book.sheetnames
    assert all(s.sheet_state == 'visible' for s in advanced)
    assert not any(suffix in str(path) for suffix in ['.csv', '.zip'])
print(f'PASS: ZIP CRC, {len(names)} XML/relationship parts, {len(book.sheetnames)} sheets, {checked} typed cells, Russian Unicode, UUID joins, styles, dates, nulls and literal formula-like text (openpyxl {openpyxl.__version__}).')
